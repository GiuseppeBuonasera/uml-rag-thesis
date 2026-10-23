"""
Analisi statistica del passo 3b (schema dei relatori, voci 156 e 173). SOLO il 3b (voce 172): i controlli oracle sono
esclusi.

Schema (dai relatori):
  1. tabella per esercizio: difficolta' 1 (Flesch-Kincaid della traccia), difficolta' 2 (size = classi + relazioni del
     diagramma di riferimento), poi una colonna di score per prompt;
  2. Spearman tra ciascuna difficolta' e ciascuna colonna di score;
  3. Wilcoxon (con t-test appaiato come confronto) per ciascuna coppia di colonne.

Scelte (voce 173):
  - colonne = 6 prompt x 2 formati = 12 (zero_shot, static k2, random k3, bm25 k2, bm25 k3, hybrid k3; PlantUML e JSON
    compatto); i relatori scrivono "8 colonne", che non corrisponde a nessuna combinazione del 3b;
  - score per esercizio = media delle ripetizioni VALIDE (non valide escluse), metrica:
      f1          F1 medio su classi, attributi, relazioni e generalizzazioni (evaluation/element_metrics.py, rilassato)
                  [principale: lo strumento dei relatori non calcola un F1];
      completezza completezza totale dello strumento dei relatori (evaluation_summary-2.xlsx), in percentuale;
  - esercizi: principale = 53 di validazione con tutte le colonne (AirTravel non ha static); sensibilita' = 79 (con i
    26 di sviluppo, voce 151);
  - correzione per test multipli: Holm, separatamente per famiglia (24 Spearman, 66 Wilcoxon, 66 t-test), alpha 0,05;
  - effetto: differenza media e mediana, r rank-biserial appaiato (Wilcoxon, differenze nulle escluse).

Uso:
    python experiments/stats_3b.py [--metrica f1|completezza] [--out results/1_passo3b]
"""

from __future__ import annotations

import argparse
import itertools
import statistics as st
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / p) for p in ("experiments", "generation", "retrieval", "")]

ALPHA = 0.05
FORMATS = (("P", "plantuml", "PlantUML"), ("C", "compatto", "JSON compatto"))
METRICS = {"f1": "F1 medio (4 tipologie, nostro)", "completezza": "Completezza totale (relatori, %)"}
SETS = {"principale": "53 di validazione (senza AirTravel)", "sensibilita": "79 (validazione + sviluppo, senza AirTravel)"}


def holm(pvalues: list[float]) -> list[float]:
    """p corretti con il metodo di Holm (step-down), nell'ordine di ingresso."""
    m = len(pvalues)
    order = sorted(range(m), key=lambda i: pvalues[i])
    adj, running = [0.0] * m, 0.0
    for rank, i in enumerate(order):
        running = max(running, min(1.0, (m - rank) * pvalues[i]))
        adj[i] = running
    return adj


def rank_biserial(diffs: list[float]) -> float | None:
    """r rank-biserial appaiato: (somma dei ranghi positivi − negativi) / somma totale; differenze nulle escluse."""
    from scipy.stats import rankdata
    d = [x for x in diffs if x != 0]
    if not d:
        return None
    r = rankdata([abs(x) for x in d])
    pos = sum(rk for rk, x in zip(r, d) if x > 0)
    neg = sum(rk for rk, x in zip(r, d) if x < 0)
    return (pos - neg) / (pos + neg)


def columns() -> list[tuple[str, str, str]]:
    import compare_3b as c3
    return [(f"{fl} · {c3.label(a, k)}", fn, c3.label(a, k)) for _, fn, fl in FORMATS for a, k in c3.CONDITIONS]


def scores(metric: str, gt: dict) -> dict:
    """(esercizio, formato, prompt) -> score medio sulle ripetizioni valide."""
    acc = defaultdict(list)
    if metric == "completezza":
        import pandas as pd
        import build_results as br
        df = pd.read_excel(br.SRC)
        for r in df[df["V"] == 1].itertuples():
            acc[(r.esercizio, r.formato, r.condizione_k)].append(float(r.completeness))
    else:
        import build_results as br
        import compare_3b as c3
        import review_report as rr
        for F, fn, _ in FORMATS:
            for run in (f"passo3b__{F}-corpus-G26", f"passo3b__{F}-testset-G26", f"passo3b_sviluppo__{F}-sviluppo-G26"):
                for x in rr.load_run(br.GEN / run, {}, gt):
                    if x["v"].level < 3:
                        continue
                    v = br.f1_mean(x)
                    if v is not None:
                        acc[(x["q"], fn, c3.label(x["m"]["condition"], x["m"]["k"]))].append(v)
    return {k: st.mean(v) for k, v in acc.items() if v}


def analyse(table: list[dict], cols: list[str]) -> dict:
    """table = righe {esercizio, insieme, fk, size, <colonna>: score}; Spearman e test appaiati con Holm."""
    from scipy.stats import spearmanr, ttest_rel, wilcoxon
    sp = []
    for dname, key in (("Flesch-Kincaid", "fk"), ("size", "size")):
        for c in cols:
            rows = [r for r in table if r.get(c) is not None]
            rho, p = spearmanr([r[key] for r in rows], [r[c] for r in rows])
            sp.append({"difficolta": dname, "colonna": c, "n": len(rows), "rho": float(rho), "p": float(p)})
    for x, a in zip(sp, holm([x["p"] for x in sp])):
        x["p_holm"] = a
    wx = []
    for a, b in itertools.combinations(cols, 2):
        rows = [r for r in table if r.get(a) is not None and r.get(b) is not None]
        d = [r[a] - r[b] for r in rows]
        w = wilcoxon(d) if any(d) else None
        t = ttest_rel([r[a] for r in rows], [r[b] for r in rows]) if any(d) else None
        wx.append({"a": a, "b": b, "n": len(rows), "media_diff": st.mean(d), "mediana_diff": st.median(d),
                   "vince_a": sum(x > 0 for x in d), "vince_b": sum(x < 0 for x in d),
                   "W": None if w is None else float(w.statistic), "p_wilcoxon": 1.0 if w is None else float(w.pvalue),
                   "r_rb": rank_biserial(d), "t": None if t is None else float(t.statistic),
                   "p_t": 1.0 if t is None else float(t.pvalue)})
    for x, a in zip(wx, holm([x["p_wilcoxon"] for x in wx])):
        x["p_wilcoxon_holm"] = a
    for x, a in zip(wx, holm([x["p_t"] for x in wx])):
        x["p_t_holm"] = a
    return {"spearman": sp, "wilcoxon": wx}


def build_table(metric: str, which: str) -> tuple[list[dict], list[str]]:
    import build_results as br
    rows, gt = br.exercise_rows()
    sc = scores(metric, gt)
    cols = columns()
    ids = [r for i, r in rows.items() if i != "AirTravel" and (which == "sensibilita" or r[1] == "valutazione")]
    table = []
    for e, ins, fk, size in br.ordered(ids, "grandezza"):
        row = {"esercizio": e, "insieme": ins, "fk": fk, "size": size}
        for name, fn, cond in cols:
            row[name] = sc.get((e, fn, cond))
        table.append(row)
    return table, [c[0] for c in cols]


def write_xlsx(out: Path, metric: str, results: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    HB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FN, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    DARK = PatternFill("solid", start_color="2F5496")
    SIG = PatternFill("solid", start_color="FFF2CC")
    nf = "0.000" if metric == "f1" else "0.00"

    def sheet(title, heads, rows, fmts, sig_col=None, widths=None):
        ws = wb.create_sheet(title)
        for j, h in enumerate(heads, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill = HB, DARK
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, r in enumerate(rows, 2):
            sig = sig_col is not None and r[sig_col] is not None and r[sig_col] < ALPHA
            for j, (v, f) in enumerate(zip(r, fmts), 1):
                c = ws.cell(row=i, column=j, value=round(v, 6) if isinstance(v, float) else v)
                c.font = BOLD if sig else FN
                if f:
                    c.number_format = f
                if sig:
                    c.fill = SIG
        for j, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(rows) + 1}"

    for which, (table, cols, res) in results.items():
        tag = "" if which == "principale" else " (79)"
        sheet(f"Tabella{tag}", ["Esercizio", "Insieme", "Difficoltà 1: Flesch-Kincaid", "Difficoltà 2: size"] + cols,
              [[r["esercizio"], r["insieme"], r["fk"], r["size"]] + [r[c] for c in cols] for r in table],
              [None, None, "0.00", "0"] + [nf] * len(cols), widths=[36, 12, 14, 12] + [13] * len(cols))
        sheet(f"Spearman{tag}", ["Difficoltà", "Colonna (formato · prompt)", "n", "rho", "p", "p Holm (24 test)"],
              [[x["difficolta"], x["colonna"], x["n"], x["rho"], x["p"], x["p_holm"]] for x in res["spearman"]],
              [None, None, "0", "+0.000;-0.000", "0.0000", "0.0000"], sig_col=5, widths=[16, 34, 6, 10, 10, 14])
        sheet(f"Wilcoxon{tag}", ["Colonna A", "Colonna B", "n", "Media A − B", "Mediana A − B", "A meglio", "B meglio",
                                 "W", "p Wilcoxon", "p Wilcoxon Holm (66)", "r rank-biserial", "t", "p t-test",
                                 "p t-test Holm (66)"],
              [[x["a"], x["b"], x["n"], x["media_diff"], x["mediana_diff"], x["vince_a"], x["vince_b"], x["W"],
                x["p_wilcoxon"], x["p_wilcoxon_holm"], x["r_rb"], x["t"], x["p_t"], x["p_t_holm"]]
               for x in res["wilcoxon"]],
              [None, None, "0", "+0.000;-0.000", "+0.000;-0.000", "0", "0", "0.0", "0.0000", "0.0000",
               "+0.00;-0.00", "+0.00;-0.00", "0.0000", "0.0000"], sig_col=9,
              widths=[30, 30, 6, 12, 12, 9, 9, 9, 11, 13, 11, 8, 11, 13])
    n = wb.create_sheet("Note")
    n.column_dimensions["A"].width = 150
    lines = [
        f"Analisi statistica del passo 3b — metrica: {METRICS[metric]}. Schema dei relatori (voci 156 e 173); solo il 3b (voce 172).",
        "  Fogli 'Tabella', 'Spearman', 'Wilcoxon': analisi principale sui 53 esercizi di validazione con tutte le colonne (AirTravel escluso: non ha static k2). Fogli '(79)': sensibilità con i 26 di sviluppo (su cui è stata tarata la configurazione: bm25 e hybrid lì sono favoriti, voce 151).",
        "  Colonne: 6 prompt × 2 formati = 12. Score per esercizio = media delle ripetizioni valide (3 per prompt; le risposte non valide sono escluse).",
        "  F1 = media degli F1 di classi, attributi, relazioni e generalizzazioni (evaluation/element_metrics.py, abbinamento rilassato). Completezza = completeness dello strumento dei relatori (recall, %).",
        "  Difficoltà 1: Flesch-Kincaid Grade Level della traccia (textstat 0.7.4). Difficoltà 2: size = classi + relazioni del diagramma di riferimento (generalizzazioni comprese). Righe ordinate per size decrescente.",
        "  Spearman: rho e p per ciascuna difficoltà e ciascuna colonna (24 test); p Holm = corretto con Holm sui 24 test.",
        "  Wilcoxon: test dei ranghi con segno appaiato per esercizio, per ciascuna coppia di colonne (66 coppie); p Holm sui 66 test. r rank-biserial = effetto (−1..+1, positivo = A meglio). t-test appaiato riportato come alternativa, con la sua correzione di Holm.",
        f"  In giallo e grassetto: risultati significativi dopo la correzione di Holm (p Holm < {ALPHA}).",
        "  Analisi a posteriori: i risultati del 3b erano già stati visti (voce 154); il disegno di questo script è registrato nella voce 173.",
        "  Generato da experiments/stats_3b.py.",
    ]
    for i, t in enumerate(lines, 1):
        c = n.cell(row=i, column=1, value=t)
        c.font = BOLD if i == 1 else FN
        c.alignment = Alignment(wrap_text=True, vertical="top")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def summary(metric: str, results: dict) -> str:
    L = [f"# Analisi statistica del passo 3b — {METRICS[metric]}", "",
         "Schema dei relatori (voci 156, 173); solo il 3b (voce 172). Correzione di Holm per famiglia, alpha 0,05.", ""]
    for which, (table, cols, res) in results.items():
        L += [f"## {SETS[which]} ({len(table)} esercizi)", "", "### Spearman (significativi dopo Holm)", ""]
        sig = [x for x in res["spearman"] if x["p_holm"] < ALPHA]
        L += [f"- {x['difficolta']} × {x['colonna']}: rho {x['rho']:+.2f}, p {x['p']:.4f}, p Holm {x['p_holm']:.4f}"
              for x in sig] or ["- nessuno"]
        L += ["", "### Wilcoxon (significativi dopo Holm)", ""]
        sig = [x for x in res["wilcoxon"] if x["p_wilcoxon_holm"] < ALPHA]
        L += [f"- {x['a']} − {x['b']}: media {x['media_diff']:+.3f}, r {x['r_rb']:+.2f}, p {x['p_wilcoxon']:.4f}, "
              f"p Holm {x['p_wilcoxon_holm']:.4f}" for x in sig] or ["- nessuno"]
        L.append("")
    return "\n".join(L)


def main(argv=None) -> int:
    a = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    a.add_argument("--metrica", choices=sorted(METRICS), default="f1")
    a.add_argument("--out", default=str(ROOT / "results" / "1_passo3b"))
    args = a.parse_args(argv)
    results = {}
    for which in SETS:
        table, cols = build_table(args.metrica, which)
        if not table or all(r[c] is None for r in table for c in cols):
            print("nessun dato del 3b: nessuna analisi scritta")
            return 1
        complete = [r for r in table if all(r[c] is not None for c in cols)] if which == "principale" else table
        results[which] = (complete, cols, analyse(complete, cols))
    out = Path(args.out)
    write_xlsx(out / f"analisi_statistica_{args.metrica}.xlsx", args.metrica, results)
    text = summary(args.metrica, results)
    md = ROOT / "data" / "results" / "generation" / "stats_3b" / f"summary_{args.metrica}.md"
    md.parent.mkdir(parents=True, exist_ok=True)
    md.write_text(text, encoding="utf-8")
    run_log = sys.modules.get("run_log")
    if run_log is not None:
        run_log.set_output(md)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(text)
    print(f"scritto in {out / f'analisi_statistica_{args.metrica}.xlsx'} e {md}")
    return 0


if __name__ == "__main__":
    import run_log  # registro delle esecuzioni (voce 102)
    with run_log.logged(__file__):
        sys.exit(main())
