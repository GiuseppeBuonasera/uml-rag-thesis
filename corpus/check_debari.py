"""
FASE 4 test set De Bari (2026-10-02): controllo indipendente del ground truth trascritto
(corpus/raw/debari_test/<id>/plantuml.txt, con le correzioni ATTIVE di corpus/corrections/)
contro docs/dati/debari/Analysis.xlsx, il materiale con cui De Bari et al. hanno valutato.

Per ogni esercizio trascritto N:
- foglio "Part 2 - N" (letto per NOME: nel file i fogli 16 e 17 sono in ordine invertito),
  colonna "Given Solution": classi, membri ("Classe .membro", metodi senza parametri) e relazioni
  ("Association (A - B)", "Generalization (Padre - Figlio)", ...), confrontati con il PlantUML
  riparsato da apollon_convert.parse_plantuml. Normalizzazione dei nomi: minuscole, solo
  caratteri alfanumerici (spazi, '-', '_', '/' e "()" ignorati).
- foglio "Estimated Difficulty": Classes / Attributes+Operations / Associations (valori in cache,
  le somme tipo '=10+7' gia' valutate) contro i conteggi del ground truth; AVG ED contro la media
  di ED 1-3.

Le discrepanze NON si correggono mai automaticamente: ognuna deve comparire in
corpus/check_debari_justifications.yaml con una classificazione
  errore_trascrizione | imprecisione_xlsx | convenzione | refuso_immagine
e una motivazione con riferimento all'immagine. Esito OK solo se ogni discrepanza e' giustificata
e ogni giustificazione corrisponde a una discrepanza reale (niente giustificazioni orfane).
Report in corpus/check_debari_report.md.

Uso:
    python corpus/check_debari.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import yaml

sys.path.insert(0, str(Path(__file__).parent))
import apollon_convert as ac
import apply_corrections as ac_corr
import build_manifest as bm

JUSTIFICATIONS_PATH = Path(__file__).parent / "check_debari_justifications.yaml"
REPORT_PATH = Path(__file__).parent / "check_debari_report.md"
CLASSIFICATIONS = {"errore_trascrizione", "imprecisione_xlsx", "convenzione", "refuso_immagine"}

REL_PREFIXES = [
    (re.compile(r"^generali", re.I), "generalization"),
    (re.compile(r"^association", re.I), "association"),
    (re.compile(r"^aggregation", re.I), "aggregation"),
    (re.compile(r"^composition", re.I), "composition"),
    (re.compile(r"^dependenc", re.I), "dependency"),
    (re.compile(r"^(implementation|realization)", re.I), "implementation"),
]
EDGE_KIND = {
    "ClassInheritance": "generalization",
    "ClassRealization": "implementation",
    "ClassAggregation": "aggregation",
    "ClassComposition": "composition",
    "ClassDependency": "dependency",
    "ClassUnidirectional": "association",
    "ClassBidirectional": "association",
}


def norm(name: str) -> str:
    return re.sub(r"[^0-9a-z]", "", name.split("(")[0].lower())


def xlsx_solution(wb, n: int) -> dict:
    """Classi, membri {(classe, membro)} e relazioni {(tipo, frozenset{A, B})} della Given Solution."""
    ws = wb[f"Part 2 - {n}"]
    header = [c.value for c in ws[2]]
    if "Given Solution" not in header:
        raise ValueError(f"Part 2 - {n}: colonna 'Given Solution' non trovata nella riga 2")
    col = header.index("Given Solution")
    classes, members, relations = set(), set(), set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        value = row[col]
        if value is None or str(value).strip() in ("", "\\"):
            continue
        value = str(value).strip()
        kind = next((k for rx, k in REL_PREFIXES if rx.match(value)), None)
        if kind:
            # "Association (Customer - Package) 1" (es. 11): numero dopo la parentesi per distinguere
            # due relazioni sulla stessa coppia
            m = re.search(r"\(([^()]*)\)", value)
            ends = [e for e in re.split(r"\s+-\s+", m.group(1))] if m else []
            if len(ends) != 2:
                raise ValueError(f"Part 2 - {n}: relazione non interpretabile: {value!r}")
            relations.add((kind, frozenset(norm(e) for e in ends)))
        elif re.match(r"^[^.(]+?\s*\.\s*\S", value):
            # "Classe .membro", anche con il punto attaccato alla classe ("Client. Name", es. 14)
            cls, member = re.split(r"\s*\.\s*", value, maxsplit=1)
            members.add((norm(cls), norm(member)))
        else:
            classes.add(norm(value))
    return {"classes": classes, "members": members, "relations": relations}


def ground_truth(model_dir: Path) -> dict:
    text = (model_dir / "plantuml.txt").read_text(encoding="utf-8")
    text, _ = ac_corr.apply_corrections(model_dir.name, text, ac_corr.load_corrections(model_dir.name))
    classes, relationships, _, unsupported = ac.parse_plantuml(text)
    if unsupported:
        raise ValueError(f"{model_dir.name}: costrutti non supportati {unsupported}")
    declared = {name: pc for name, pc in classes.items() if not pc.placeholder}
    members = set()
    for name, pc in declared.items():
        if pc.kind == "enum":
            continue
        members |= {(norm(name), norm(a)) for a, _ in pc.attributes}
        members |= {(norm(name), norm(m.lstrip("+-#~ "))) for m in pc.methods}
    relations = set()
    for r in relationships:
        if r["kind"] == "binary":
            edge_type, _, _ = ac.relationship_kind(r["op"])
            relations.add((EDGE_KIND[edge_type], frozenset({norm(r["source"]), norm(r["target"])})))
        else:
            relations.add(("association_class", frozenset({norm(r["a"]), norm(r["b"])})))
    n_members = sum(len(pc.attributes) + len(pc.methods) for pc in declared.values() if pc.kind != "enum")
    return {
        "classes": {norm(c) for c in declared},
        "members": members,
        "relations": relations,
        "counts": {"classes": len(declared), "attributes_operations": n_members,
                   "associations": len(relationships)},
    }


def fmt_rel(rel) -> str:
    kind, ends = rel
    return f"{kind}({' - '.join(sorted(ends))})"


def discrepancies(n: int, gt: dict, xs: dict, difficulty: dict) -> list[str]:
    out = []
    out += [f"classe solo nell'xlsx: {c}" for c in sorted(xs["classes"] - gt["classes"])]
    out += [f"classe solo nel ground truth: {c}" for c in sorted(gt["classes"] - xs["classes"])]
    out += [f"membro solo nell'xlsx: {c}.{m}" for c, m in sorted(xs["members"] - gt["members"])]
    out += [f"membro solo nel ground truth: {c}.{m}" for c, m in sorted(gt["members"] - xs["members"])]
    out += [f"relazione solo nell'xlsx: {fmt_rel(r)}" for r in sorted(xs["relations"] - gt["relations"], key=fmt_rel)]
    out += [f"relazione solo nel ground truth: {fmt_rel(r)}"
            for r in sorted(gt["relations"] - xs["relations"], key=fmt_rel)]
    for key, label in (("classes", "Classes"), ("attributes_operations", "Attributes + Operations"),
                       ("associations", "Associations")):
        if difficulty["debari_xlsx_counts"][key] != gt["counts"][key]:
            out.append(f"conteggio {label}: xlsx {difficulty['debari_xlsx_counts'][key]}, "
                       f"ground truth {gt['counts'][key]}")
    return out


def main() -> None:
    wb = openpyxl.load_workbook(bm.DEBARI_XLSX, data_only=True)
    difficulty = bm.debari_difficulty()
    ws_ed = wb["Estimated Difficulty"]
    cached_avg = {int(r[1]): r[8] for r in ws_ed.iter_rows(values_only=True) if isinstance(r[1], (int, float))}
    justifications = yaml.safe_load(JUSTIFICATIONS_PATH.read_text(encoding="utf-8")) if JUSTIFICATIONS_PATH.exists() else {}
    justifications = justifications or {}

    lines = ["# Controllo test set De Bari contro Analysis.xlsx", "",
             "Generato da `corpus/check_debari.py` (non modificare a mano).", ""]
    unjustified, orphans, done = [], [], []
    for model_dir in bm.list_model_dirs([bm.DEBARI_RAW_DIR]):
        n = int(bm.DEBARI_ID_RE.match(model_dir.name).group(1))
        if not (model_dir / "plantuml.txt").exists():
            continue
        done.append(model_dir.name)
        gt = ground_truth(model_dir)
        found = discrepancies(n, gt, xlsx_solution(wb, n), difficulty[n])
        if abs(cached_avg[n] - difficulty[n]["debari_ed_avg"]) > 1e-3:
            found.append(f"AVG ED in cache {cached_avg[n]} diverso dalla media di ED 1-3 {difficulty[n]['debari_ed_avg']}")
        given = {j["discrepanza"]: j for j in justifications.get(model_dir.name, [])}
        for j in given.values():
            if j.get("classificazione") not in CLASSIFICATIONS:
                raise ValueError(f"{model_dir.name}: classificazione non ammessa {j.get('classificazione')!r}")
        lines += [f"## {model_dir.name}", "",
                  f"Conteggi ground truth: {gt['counts']} — xlsx: {difficulty[n]['debari_xlsx_counts']}; "
                  f"ED medio {difficulty[n]['debari_ed_avg']}", ""]
        if not found:
            lines += ["Nessuna discrepanza.", ""]
        for d in found:
            j = given.get(d)
            if j:
                lines.append(f"- {d} — **{j['classificazione']}**: {j['motivo']}")
            else:
                lines.append(f"- {d} — **NON GIUSTIFICATA**")
                unjustified.append(f"{model_dir.name}: {d}")
        orphans += [f"{model_dir.name}: {d}" for d in given if d not in found]
        lines.append("")

    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")
    print(f"Esercizi controllati: {len(done)} {done}")
    print(f"Report: {REPORT_PATH}")
    if orphans:
        print("Giustificazioni senza discrepanza corrispondente:", *orphans, sep="\n  ")
    if unjustified:
        print(f"{len(unjustified)} discrepanze NON giustificate:", *unjustified, sep="\n  ")
    if unjustified or orphans:
        raise SystemExit(1)
    print("check_debari OK: ogni discrepanza e' giustificata.")


if __name__ == "__main__":
    main()
